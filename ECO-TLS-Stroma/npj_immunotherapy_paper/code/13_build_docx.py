"""Combine manuscript parts and rebuild DOCX (9 figs, 4 tables)."""
import re
from docx import Document
from docx.shared import Pt, Inches, RGBColor

BASE = "/home/user/npj_immunotherapy_paper"
MAN, FIG = f"{BASE}/manuscript", f"{BASE}/results/figures"
TITLE = ("A TLS-versus-Stroma Ecosystem Score Predicts Checkpoint Inhibitor Benefit "
         "Across Melanoma, NSCLC and Bladder Cancer")

p1 = open(f"{MAN}/manuscript_part1.md").read().rstrip() + "\n"
p2 = open(f"{MAN}/manuscript_part2.md").read()
md = p1 + p2
open(f"{MAN}/manuscript.md", "w").write(md)
print("manuscript.md words:", len(md.split()), flush=True)

doc = Document()
style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(10.5)
for i in (1, 2, 3):
    doc.styles[f"Heading {i}"].font.color.rgb = RGBColor(0x1F, 0x3B, 0x63)
doc.add_heading(TITLE, level=1)
doc.add_paragraph("[First Author]¹, [Second Author]², [Corresponding Author]¹* — ¹[Affiliation 1]; ²[Affiliation 2] — *Corresponding: [email]")
doc.add_paragraph("Article type: Original Research Article — npj Precision Oncology Collection: Next-Generation Precision Immunotherapy",
                  style="Intense Quote")
for line in md.split("\n"):
    if line.startswith("# "):
        continue
    if line.startswith("**Authors:**") or line.startswith("**Affiliations:**") or line.startswith("**Article type:**") or line.startswith("> **Author note:**"):
        continue
    if line.startswith("## "):
        doc.add_heading(line[3:], level=2)
        continue
    if line.startswith("### "):
        doc.add_heading(line[4:], level=3)
        continue
    if line.startswith("|"):
        continue
    if line.strip() == "" or line.strip() == "---":
        continue
    p = doc.add_paragraph()
    for seg in re.split(r"(\*\*.+?\*\*)", line):
        if seg.startswith("**") and seg.endswith("**") and len(seg) > 4:
            r = p.add_run(seg[2:-2]); r.bold = True
        else:
            p.add_run(seg)

def add_table(title, headers, rows):
    doc.add_heading(title, level=3)
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Light Grid Accent 1"
    for k, h in enumerate(headers):
        t.rows[0].cells[k].text = h
    for i, row in enumerate(rows, start=1):
        for k, val in enumerate(row):
            t.rows[i].cells[k].text = str(val)

add_table("Table 1. Cohort summary",
          ["Cohort", "Cancer", "Setting", "n (analysis)", "Endpoint", "Source"],
          [["TCGA-SKCM", "Melanoma", "Prognostic", "440", "OS (153 events)", "Firehose"],
           ["TCGA-LUAD", "NSCLC (adeno)", "Prognostic", "477", "OS (121 events)", "Firehose"],
           ["TCGA-BLCA", "Bladder", "Prognostic", "394", "OS (107 events)", "Firehose"],
           ["GSE78220", "Melanoma", "anti-PD-1", "25 (13 R)", "RECIST", "GEO/Hugo 2016"],
           ["GSE91061", "Melanoma", "Nivolumab (pre)", "49 (10 R)", "RECIST", "GEO/Riaz 2017"],
           ["GSE126044", "NSCLC", "anti-PD-1", "16 (5 R)", "Response", "GEO"],
           ["GSE135222", "NSCLC", "anti-PD-1/PD-L1", "27 (7 DCB)", "DCB + PFS (21 ev)", "GEO"],
           ["GSE207422", "NSCLC", "anti-PD-1+chemo neoadj (pre)", "24 (9 MPR)", "MPR", "GEO"],
           ["IMvigor210", "Bladder", "Atezolizumab", "298 R / 326 OS (213 ev)", "mRECIST + OS + TMB", "cBioPortal"]])
add_table("Table 2. Response discrimination (AUC, 95% bootstrap CI for ECO)",
          ["Cohort", "ECO", "TLS", "STROMA", "IFNG6", "CD8EFF", "CD274"],
          [["GSE78220", "0.71 [0.47-0.90]", "0.45", "0.21", "0.46", "0.47", "0.56"],
           ["GSE91061", "0.60 [0.39-0.81]", "0.66", "0.49", "0.63", "0.62", "0.54"],
           ["GSE126044", "0.78 [0.44-1.00]", "0.69", "0.56", "0.75", "0.87", "0.60"],
           ["GSE135222", "0.64 [0.35-0.90]", "0.53", "0.33", "0.73", "0.74", "0.61"],
           ["GSE207422", "0.75 [0.52-0.93]", "0.74", "0.40", "0.63", "0.76", "0.61"],
           ["IMvigor210", "0.60 [0.53-0.67]", "0.54", "0.43", "0.61", "0.61", "0.57"],
           ["Pooled (n=439)", "0.62", "—", "—", "—", "—", "—"]])
add_table("Table 3. Survival by ECO-high vs ECO-low",
          ["Cohort", "Endpoint", "n (events)", "HR [95% CI]", "p", "Adj. HR", "Adj. p"],
          [["TCGA-SKCM", "OS", "440 (153)", "0.53 [0.38-0.73]", "0.0001", "0.60", "0.002"],
           ["TCGA-LUAD", "OS", "477 (121)", "0.63 [0.44-0.90]", "0.011", "0.69", "0.052"],
           ["TCGA-BLCA", "OS", "394 (107)", "0.64 [0.44-0.95]", "0.025", "0.76", "0.17"],
           ["IMvigor210", "OS", "326 (213)", "0.69 [0.52-0.90]", "0.006", "0.84", "0.28"],
           ["GSE135222", "PFS", "27 (21)", "0.53 [0.22-1.27]", "0.15", "—", "—"]])
add_table("Table 4. Treatment ideas for ECO-low/stromal patients",
          ["Hub gene(s)", "Program", "Drug class", "Proposed concept"],
          [["TGFB1, TGFBR2", "Stroma", "Pirfenidone; PD-L1xTGF-b bispecifics", "Checkpoint + TGF-b blockade in TGFB/COL-high"],
           ["CXCL12, CXCR4", "Bridge", "Plerixafor, motixafortide, mavorixafor", "Checkpoint + CXCR4 blocker in CXCL12-high"],
           ["VEGFA", "Stroma/vessels", "Bevacizumab, aflibercept, pazopanib", "Checkpoint + VEGF blockade"],
           ["FAP", "CAF", "FAP-directed (CAR, IL2v, radioligands; exp.)", "Checkpoint + FAP targeting in FAP-heavy CAF"],
           ["ENTPD1, NT5E (CD39, CD73)", "Adenosine", "CD39, CD73 blockers (exp.)", "Checkpoint + adenosine blockade in NT5E-high"],
           ["IDO1", "Metabolic", "Epacadostat (caution: Ph3 fail)", "Only in biomarker-selected re-testing"],
           ["POSTN, MMP14, collagens", "Matrix", "Antifibrotics; miR-29 mimics (preclin.)", "Stromal softening + checkpoint"],
           ["CD274, CTLA4, LAG3", "TLS/effector", "Approved checkpoint drugs", "ECO-high: checkpoint therapy per label"]])

doc.add_heading("Figures (embedded for review; high-resolution files submitted separately)", level=2)
for i, cap in enumerate(["Fig1_design", "Fig2_response", "Fig3_benchmark", "Fig4_survival", "Fig5_subtypes",
                         "Fig6_TMB", "Fig7_mechanism", "Fig8_regulation_drugs", "Fig9_translation"], start=1):
    doc.add_paragraph(f"Figure {i}", style="Heading 3")
    doc.add_picture(f"{FIG}/{cap}.png", width=Inches(6.2))
doc.save(f"{MAN}/manuscript.docx")

from docx import Document as D2
d = D2(f"{MAN}/manuscript.docx")
words = sum(len(p.text.split()) for p in d.paragraphs)
print("DOCX paragraphs:", len(d.paragraphs), "| approx words:", words, "| tables:", len(d.tables), flush=True)
# reference count check
import re as _re
refs = _re.findall(r"^(\d+)\.\s+\S", md, flags=_re.M)
print("References found:", len(refs), "| last number:", refs[-1] if refs else None, flush=True)
